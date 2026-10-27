from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path, PurePosixPath
import sqlite3
import sys
import uuid
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from app_runtime import storage_directory  # noqa: E402
from backend import database  # noqa: E402
from backend.services.common import object_status  # noqa: E402
from simulation_store import simulation_guard, simulation_path, validate_simulation  # noqa: E402

WORKBOOK = ROOT_DIR / "pictures" / "sptest.xlsx"
IMAGE_ARCHIVE = ROOT_DIR / "pictures" / "2587aecf-8770-458a-a961-3a13aedd5ee3.zip"
CATEGORY = "Okategoriserad"
CORRECTED_ROW_VOLVO_IDS = {42: "004712", 77: "004713"}


def clean_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def header_key(value: object) -> str:
    text = clean_text(value).casefold()
    return (text.replace("å", "a").replace("ä", "a").replace("ö", "o")
            .replace(" ", "_").replace("/", "_").replace("-", "_"))


def parse_price(value: object) -> str | None:
    text = clean_text(value)
    if not text:
        return None
    text = text.replace("\u00a0", "").replace(" ", "").replace("kr", "").replace("KR", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        number = Decimal(text)
    except InvalidOperation as error:
        raise ValueError(f"Ogiltigt styckpris: {value!r}") from error
    if not number.is_finite() or number < 0:
        raise ValueError(f"Ogiltigt styckpris: {value!r}")
    return format(number.quantize(Decimal("0.01")), "f")


def parse_unit_ids(value: object) -> list[str]:
    lines = [line.strip() for line in clean_text(value).splitlines()]
    return list(dict.fromkeys(line for line in lines if line and line.casefold() != "individnr"))


def image_filename(value: object) -> str:
    raw = clean_text(value)
    if not raw:
        return ""
    for field in ("fileName", "serverRelativeUrl"):
        marker = f'"{field}":"'
        if marker in raw:
            path = raw.split(marker, 1)[1].split('"', 1)[0].replace("\\/", "/")
            return PurePosixPath(path).name
        marker = f'"{field}": "'
        if marker in raw:
            path = raw.split(marker, 1)[1].split('"', 1)[0].replace("\\/", "/")
            return PurePosixPath(path).name
    return ""


def load_rows(workbook_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [header_key(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(clean_text(value) for value in values):
            continue
        rows.append({**dict(zip(headers, values)), "_excel_row": excel_row})
    return rows


def image_index(archive: ZipFile) -> tuple[dict[str, str], set[str]]:
    candidates: dict[str, list[str]] = defaultdict(list)
    for member in archive.namelist():
        if member.endswith("/"):
            continue
        candidates[PurePosixPath(member.replace("\\", "/")).name.casefold()].append(member)
    unique = {name: members[0] for name, members in candidates.items() if len(members) == 1}
    ambiguous = {name for name, members in candidates.items() if len(members) > 1}
    return unique, ambiguous


def image_data(archive: ZipFile, member: str) -> str:
    from PIL import Image, ImageOps

    with Image.open(BytesIO(archive.read(member))) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        image.thumbnail((900, 900))
        for quality in (82, 74, 66, 58, 50, 44):
            buffer = BytesIO()
            image.save(buffer, format="JPEG", quality=quality, optimize=True)
            encoded = buffer.getvalue()
            data_url = "data:image/jpeg;base64," + __import__("base64").b64encode(encoded).decode("ascii")
            if len(data_url) <= 700_000:
                return data_url
    raise ValueError(f"Bilden kunde inte skalas till godkänd storlek: {member}")


def build_import_rows(rows: list[dict[str, object]], archive: ZipFile) -> tuple[list[dict[str, object]], dict[str, int]]:
    by_filename, ambiguous = image_index(archive)
    duplicate_ids: defaultdict[str, list[int]] = defaultdict(list)
    parsed_rows = []
    for row in rows:
        excel_row = int(row["_excel_row"])
        unit_ids = parse_unit_ids(row.get("id_nummer"))
        for unit_id in unit_ids:
            duplicate_ids[unit_id].append(excel_row)
        filename = image_filename(row.get("bild"))
        image_member = by_filename.get(filename.casefold()) if filename else None
        parsed_rows.append({
            "excel_row": excel_row,
            "unit_ids": unit_ids,
            "image_member": image_member,
            "image_filename": filename,
            "row": row,
        })

    duplicate_groups = {value: locations for value, locations in duplicate_ids.items() if len(set(locations)) > 1}
    unexpected_duplicates = {
        value: locations for value, locations in duplicate_groups.items()
        if not set(locations).issubset(CORRECTED_ROW_VOLVO_IDS)
    }
    if unexpected_duplicates:
        samples = list(unexpected_duplicates.items())[:5]
        raise ValueError(f"ID:n delas mellan ovantade rader: {samples}")

    output = []
    image_count = 0
    for item in parsed_rows:
        excel_row = item["excel_row"]
        source = item["row"]
        unit_ids = item["unit_ids"]
        corrected_volvo_id = CORRECTED_ROW_VOLVO_IDS.get(excel_row, "")
        correction = bool(corrected_volvo_id)
        quantity = 1 if correction else len(unit_ids) or 1
        notes = [f"SharePoint Excelrad {excel_row}"]
        if correction:
            notes.append(f"Delad källista korrigerad; tilldelat VOLVO-ID {corrected_volvo_id}.")
        elif not unit_ids:
            notes.append("Saknar ID-nummer; testsaldo satt till 1.")
        if item["image_filename"] and not item["image_member"]:
            notes.append(f"Bild saknas i ZIP: {item['image_filename']}")
        if unit_ids and not correction:
            notes.append("VOLVO-ID: " + ", ".join(unit_ids))

        encoded_image = ""
        if item["image_member"]:
            encoded_image = image_data(archive, item["image_member"])
            image_count += 1

        price = parse_price(source.get("price_per_unit"))
        object_id = f"SP-{excel_row:04d}"
        output.append({
            "object_id": object_id,
            "name": clean_text(source.get("title")) or object_id,
            "category": CATEGORY,
            "brand": clean_text(source.get("fabrikat")),
            "model": clean_text(source.get("modell")),
            "measurement_range": clean_text(source.get("matomrade")),
            "unit_price": price,
            "article_number": "",
            "unit_volvo_ids": corrected_volvo_id or "\n".join(unit_ids),
            "item_type": clean_text(source.get("item_type")),
            "description": "",
            "location": clean_text(source.get("placering")) or "Placering saknas",
            "storage_location_2": "",
            "tornado_id": "",
            "door_id": "",
            "belongs_to_instrument": "",
            "device": "",
            "belongs_to": "",
            "last_changed": "",
            "search_words": "; ".join(notes),
            "tray_for": "",
            "lift_tray": "",
            "total_quantity": quantity,
            "available_quantity": quantity,
            "min_quantity": 0,
            "status": object_status(quantity, 0),
            "is_consumable": 0,
            "image_data": encoded_image or None,
        })

    summary = {
        "rows": len(output),
        "quantity": sum(int(item["total_quantity"]) for item in output),
        "images": image_count,
        "unmatched_images": sum(bool(item["image_filename"]) and not item["image_member"] for item in parsed_rows),
        "rows_without_ids": sum(not item["unit_ids"] and item["excel_row"] not in CORRECTED_ROW_VOLVO_IDS for item in parsed_rows),
        "correction_rows": len(CORRECTED_ROW_VOLVO_IDS),
        "ambiguous_archive_names": len(ambiguous),
        "prices": sum(item["unit_price"] is not None for item in output),
    }
    return output, summary


def backup_database(path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = path.with_name(f"demo.before-sharepoint-{stamp}-{uuid.uuid4().hex[:8]}.db")
    with sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True) as source:
        with sqlite3.connect(target) as destination:
            source.backup(destination)
            if destination.execute("PRAGMA quick_check").fetchone()[0] != "ok":
                target.unlink(missing_ok=True)
                raise ValueError("Säkerhetskopian klarade inte SQLite-kontrollen.")
    return target


def print_summary(summary: dict[str, int], rows: list[dict[str, object]], path: Path, apply: bool) -> None:
    print(f"Simuleringsdatabas: {path}")
    print(f"SharePoint-rader: {summary['rows']}")
    print(f"Planerat saldo: {summary['quantity']}")
    print(f"Bilder matchade exakt: {summary['images']}")
    print(f"Bildreferenser utan matchning: {summary['unmatched_images']}")
    print(f"Rader utan ID (tillfälligt saldo 1): {summary['rows_without_ids']}")
    print(f"Korrigerade rader med saldo 1: {summary['correction_rows']}")
    print(f"Styckpriser: {summary['prices']}")
    for row_number in (42, 77):
        item = next((value for value in rows if value["object_id"] == f"SP-{row_number:04d}"), None)
        if item:
            print(f"  {item['object_id']} | {item['name']} | saldo {item['total_quantity']} | VOLVO-ID {item['unit_volvo_ids']}")
    print("Läge: SKRIVER" if apply else "Läge: TORRKÖRNING - inget har ändrats")


def repair_corrected_rows(target: Path, apply: bool) -> None:
    expected = (("SP-0042", "Kabel", "004712"), ("SP-0077", "Termineringsplugg", "004713"))
    with sqlite3.connect(target.resolve().as_uri() + "?mode=ro", uri=True) as db:
        for object_id, name, volvo_id in expected:
            item = db.execute(
                "SELECT name,total_quantity,unit_volvo_ids FROM objects WHERE object_id=?", (object_id,)
            ).fetchone()
            if not item or item[0] != name or item[1] != 1:
                raise ValueError(f"Hittade inte förväntad simuleringspost {object_id} ({name}, saldo 1).")
            if item[2] not in (None, "", volvo_id):
                raise ValueError(f"{object_id} har redan ett annat exemplars-VOLVO-ID; ingen data ändrad.")
            collision = db.execute(
                "SELECT object_id FROM objects WHERE unit_volvo_ids=? AND object_id<>? LIMIT 1",
                (volvo_id, object_id),
            ).fetchone()
            if collision:
                raise ValueError(f"VOLVO-ID {volvo_id} används redan av {collision[0]}; ingen data ändrad.")

    print("Rättelser:")
    for object_id, name, volvo_id in expected:
        print(f"  {object_id} | {name} | saldo 1 | VOLVO-ID {volvo_id}")
    print("Läge: SKRIVER" if apply else "Läge: TORRKÖRNING - inget har ändrats")
    if not apply:
        return

    backup = backup_database(target)
    print(f"Backup: {backup}")
    with database.connect() as db:
        db.execute("BEGIN IMMEDIATE")
        for object_id, name, volvo_id in expected:
            note = f"SharePoint Excelrad {object_id[-4:]}; delad källista korrigerad; tilldelat VOLVO-ID {volvo_id}."
            cursor = db.execute(
                "UPDATE objects SET unit_volvo_ids=?,search_words=? WHERE object_id=? AND name=? AND total_quantity=1",
                (volvo_id, note, object_id, name),
            )
            if cursor.rowcount != 1:
                raise ValueError(f"Kunde inte rätta {object_id}; transaktionen avbryts.")
    print("Rättelse klar; båda objekten har olika exemplar-VOLVO-ID:n.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Importera SharePoint-objekt till den sparade simuleringen.")
    parser.add_argument("--apply", action="store_true", help="Säkerhetskopiera och skriv till den sparade simuleringen.")
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--images", type=Path, default=IMAGE_ARCHIVE)
    parser.add_argument("--repair-corrected-ids", action="store_true",
                        help="Torrkör eller rätta de två redan importerade simuleringsposterna.")
    args = parser.parse_args()

    if not args.workbook.is_file() or not args.images.is_file():
        raise SystemExit("Arbetsboken eller ZIP-filen saknas i pictures/.")

    target = simulation_path(storage_directory())
    validate_simulation(target)
    database.DATABASE_PATH = target
    if args.repair_corrected_ids:
        with simulation_guard(target):
            validate_simulation(target)
            repair_corrected_rows(target, apply=args.apply)
        return

    with simulation_guard(target):
        validate_simulation(target)
        rows = load_rows(args.workbook)
        with ZipFile(args.images) as archive:
            imported, summary = build_import_rows(rows, archive)

        with sqlite3.connect(target.resolve().as_uri() + "?mode=ro", uri=True) as db:
            existing = db.execute("SELECT COUNT(*) FROM objects WHERE object_id LIKE 'SP-%'").fetchone()[0]
        if existing:
            raise SystemExit(f"Hittade {existing} befintliga SP-objekt. Avbryter för att undvika dubbelimport.")

        print_summary(summary, imported, target, args.apply)
        if not args.apply:
            return

        backup = backup_database(target)
        print(f"Backup: {backup}")
        database.initialize_database()
        with database.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            db.execute("INSERT OR IGNORE INTO object_categories(name) VALUES (?)", (CATEGORY,))
            db.executemany(
                """INSERT INTO objects (
                    object_id, name, category, brand, model, measurement_range, unit_price,
                    article_number, unit_volvo_ids, item_type, description, location,
                    storage_location_2, tornado_id, door_id, belongs_to_instrument, device,
                    belongs_to, last_changed, search_words, tray_for, lift_tray,
                    total_quantity, available_quantity, min_quantity, status, is_consumable, image_data
                ) VALUES (
                    :object_id, :name, :category, :brand, :model, :measurement_range, :unit_price,
                    :article_number, :unit_volvo_ids, :item_type, :description, :location,
                    :storage_location_2, :tornado_id, :door_id, :belongs_to_instrument, :device,
                    :belongs_to, :last_changed, :search_words, :tray_for, :lift_tray,
                    :total_quantity, :available_quantity, :min_quantity, :status, :is_consumable, :image_data
                )""",
                imported,
            )
            db.execute(
                "INSERT INTO events(event_time,event_type,message) VALUES (?, 'import', ?)",
                (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                 f"SharePoint-import till simulering: {summary['rows']} objekt."),
            )
        print(f"Import klar. Befintliga demoobjekt behölls; tillagda objekt: {summary['rows']}.")


if __name__ == "__main__":
    main()
