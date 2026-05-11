from __future__ import annotations

import argparse
import re
import shutil
import sqlite3
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from backend.database import DATABASE_PATH, connect, initialize_database  # noqa: E402
from backend.services import object_status  # noqa: E402


HEADER_ROW = 2
DATA_START_ROW = 3
IMPORT_SOURCE = "Alla artiklar och reservdelar.xlsx"


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def clean_int(value: Any, default: int = 0) -> int:
    if value is None or value == "":
        return default
    if isinstance(value, str):
        value = value.strip().replace(" ", "").replace(",", ".")
        if not value:
            return default
    try:
        return max(0, int(float(value)))
    except (TypeError, ValueError):
        return default


def normalize_header(value: Any, index: int, seen: Counter[str]) -> str:
    text = clean_text(value) or f"column_{index}"
    text = text.lower()
    text = text.replace("å", "a").replace("ä", "a").replace("ö", "o")
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    text = text or f"column_{index}"
    seen[text] += 1
    if seen[text] > 1:
        return f"{text}_{seen[text]}"
    return text


def load_excel_rows(excel_path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    sheet = workbook.active
    raw_headers = next(sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))
    seen: Counter[str] = Counter()
    headers = [normalize_header(value, index, seen) for index, value in enumerate(raw_headers, start=1)]

    rows: list[dict[str, Any]] = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=DATA_START_ROW, values_only=True), start=DATA_START_ROW):
        if not any(clean_text(value) for value in values):
            continue
        row = dict(zip(headers, values))
        row["_excel_row"] = excel_row
        rows.append(row)
    return headers, rows


def make_object_ids(rows: list[dict[str, Any]]) -> dict[int, str]:
    bases = [clean_text(row.get("volvoid_artikelnummer")) or f"IMPORT-{row['_excel_row']}" for row in rows]
    counts = Counter(bases)
    used: set[str] = set()
    duplicate_index: defaultdict[str, int] = defaultdict(int)
    object_ids: dict[int, str] = {}

    for row, base in zip(rows, bases):
        object_id = base
        if counts[base] > 1:
            duplicate_index[base] += 1
            object_id = f"{base}-{duplicate_index[base]}"

        while object_id in used:
            duplicate_index[base] += 1
            object_id = f"{base}-{duplicate_index[base]}"

        used.add(object_id)
        object_ids[row["_excel_row"]] = object_id

    return object_ids


def map_object(row: dict[str, Any], object_id: str) -> tuple[Any, ...]:
    article_number = clean_text(row.get("volvoid_artikelnummer"))
    category = clean_text(row.get("kategori")) or "Okategoriserad"
    item_type = clean_text(row.get("typ"))
    description = clean_text(row.get("beteckning"))
    supplier_article_number = clean_text(row.get("leverantorens_artikelnummer"))
    name = clean_text(row.get("namn")) or clean_text(row.get("beteckning")) or article_number or object_id
    location = clean_text(row.get("lagerplats")) or clean_text(row.get("lagerplats_2")) or "Ej angiven"
    storage_location_2 = clean_text(row.get("lagerplats_2"))
    tornado_id = clean_text(row.get("forrad")) or "Ej angivet"
    door_id = clean_text(row.get("hissbricka")) or clean_text(row.get("lagerplats_2")) or clean_text(row.get("lagerplats")) or "-"
    belongs_to_instrument = clean_text(row.get("tillhor_instrument"))
    device = clean_text(row.get("device"))
    belongs_to = clean_text(row.get("tillhor"))
    last_changed = clean_text(row.get("senast_andrad"))
    search_words = clean_text(row.get("sok_ord"))
    tray_for = clean_text(row.get("ligger_pa_bricka_for"))
    lift_tray = clean_text(row.get("hissbricka"))
    total_quantity = clean_int(row.get("saldo"))
    available_quantity = total_quantity
    min_quantity = clean_int(row.get("bestallningspunkt"))
    status = object_status(available_quantity, min_quantity)

    return (
        object_id,
        name,
        category,
        item_type,
        description,
        article_number,
        supplier_article_number,
        location,
        storage_location_2,
        tornado_id,
        door_id,
        belongs_to_instrument,
        device,
        belongs_to,
        last_changed,
        search_words,
        tray_for,
        lift_tray,
        total_quantity,
        available_quantity,
        min_quantity,
        status,
    )


def ensure_raw_import_table(db: sqlite3.Connection, headers: list[str]) -> None:
    columns = [
        "import_id INTEGER PRIMARY KEY AUTOINCREMENT",
        "source_file TEXT NOT NULL",
        "source_row INTEGER NOT NULL",
        "imported_at TEXT NOT NULL",
        "object_id TEXT NOT NULL",
    ]
    columns.extend(f'"{header}" TEXT' for header in headers)
    db.execute("DROP TABLE IF EXISTS old_inventory_import")
    db.execute(f"CREATE TABLE old_inventory_import ({', '.join(columns)})")


def insert_raw_rows(db: sqlite3.Connection, headers: list[str], rows: list[dict[str, Any]], object_ids: dict[int, str]) -> None:
    columns = ["source_file", "source_row", "imported_at", "object_id", *headers]
    placeholders = ", ".join("?" for _ in columns)
    quoted_columns = ", ".join(f'"{column}"' for column in columns)
    imported_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    values = []
    for row in rows:
        values.append(
            [
                IMPORT_SOURCE,
                row["_excel_row"],
                imported_at,
                object_ids[row["_excel_row"]],
                *[clean_text(row.get(header)) for header in headers],
            ]
        )

    db.executemany(
        f"INSERT INTO old_inventory_import ({quoted_columns}) VALUES ({placeholders})",
        values,
    )


def backup_database() -> Path | None:
    if not DATABASE_PATH.exists():
        return None
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = DATABASE_PATH.with_name(f"{DATABASE_PATH.stem}_backup_before_excel_import_{timestamp}{DATABASE_PATH.suffix}")
    shutil.copy2(DATABASE_PATH, backup_path)
    return backup_path


def import_rows(headers: list[str], rows: list[dict[str, Any]], apply: bool) -> None:
    object_ids = make_object_ids(rows)
    duplicate_count = len(rows) - len({clean_text(row.get("volvoid_artikelnummer")) for row in rows})
    preview = [map_object(row, object_ids[row["_excel_row"]]) for row in rows[:5]]

    print(f"Rader i Excel: {len(rows)}")
    print(f"Dubbletter i Volvoid/Artikelnummer: {duplicate_count}")
    print("Exempel pa importerade objects-rader:")
    for item in preview:
        print(f"  {item[0]} | {item[2]} | {item[1]} | saldo {item[18]} | plats {item[7]}")

    if not apply:
        print("Dry run klar. Inget skrevs till databasen.")
        return

    initialize_database()
    backup_path = backup_database()
    if backup_path:
        print(f"Backup skapad: {backup_path}")

    object_values = [map_object(row, object_ids[row["_excel_row"]]) for row in rows]
    with connect() as db:
        ensure_raw_import_table(db, headers)
        insert_raw_rows(db, headers, rows, object_ids)
        db.executemany(
            """
            INSERT INTO objects (
                object_id, name, category, item_type, description, article_number, supplier_article_number,
                location, storage_location_2, tornado_id, door_id, belongs_to_instrument, device,
                belongs_to, last_changed, search_words, tray_for, lift_tray,
                total_quantity, available_quantity, min_quantity, status
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(object_id) DO UPDATE SET
                name = excluded.name,
                category = excluded.category,
                item_type = excluded.item_type,
                description = excluded.description,
                article_number = excluded.article_number,
                supplier_article_number = excluded.supplier_article_number,
                location = excluded.location,
                storage_location_2 = excluded.storage_location_2,
                tornado_id = excluded.tornado_id,
                door_id = excluded.door_id,
                belongs_to_instrument = excluded.belongs_to_instrument,
                device = excluded.device,
                belongs_to = excluded.belongs_to,
                last_changed = excluded.last_changed,
                search_words = excluded.search_words,
                tray_for = excluded.tray_for,
                lift_tray = excluded.lift_tray,
                total_quantity = excluded.total_quantity,
                available_quantity = excluded.available_quantity,
                min_quantity = excluded.min_quantity,
                status = excluded.status
            """,
            object_values,
        )
        db.execute(
            "INSERT INTO events (event_time, event_type, message) VALUES (?, ?, ?)",
            (
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "import",
                f"Importerade {len(rows)} artiklar fran Excel",
            ),
        )

    print(f"Import klar: {len(rows)} rader skrevs till objects och old_inventory_import.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Importera artiklar fran Excel till SQLite.")
    parser.add_argument("excel_file", type=Path)
    parser.add_argument("--apply", action="store_true", help="Skriv importen till databasen.")
    args = parser.parse_args()

    if not args.excel_file.exists():
        raise SystemExit(f"Filen hittades inte: {args.excel_file}")

    headers, rows = load_excel_rows(args.excel_file)
    import_rows(headers, rows, apply=args.apply)


if __name__ == "__main__":
    main()
